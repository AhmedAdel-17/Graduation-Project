"""One-off: convert the literature-review markdown into a styled .docx.

Handles the subset of Markdown used in docs/literature_review_agents.md:
headings, paragraphs, pipe tables, ordered/unordered lists, blockquotes,
horizontal rules, and inline **bold** / *italic* / `code` / [text](url).
"""
import re
import sys
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

SRC = sys.argv[1]
OUT = sys.argv[2]

INLINE_RX = re.compile(
    r"(\*\*.+?\*\*|(?<!\*)\*[^*].*?\*|`[^`]+`|\[[^\]]+\]\([^)]+\))"
)
LINK_RX = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")


def add_inline(paragraph, text):
    """Render inline markdown spans into runs on an existing paragraph."""
    for part in INLINE_RX.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        elif LINK_RX.fullmatch(part):
            m = LINK_RX.fullmatch(part)
            run = paragraph.add_run(m.group(1))
            run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            paragraph.add_run(part)


def strip_inline(text):
    """Plain-text version of inline markdown (for table cells)."""
    text = LINK_RX.sub(r"\1", text)
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"(?<!\*)\*([^*]+?)\*", r"\1", text)
    return text.strip()


def parse_table(rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    header, body = cells[0], cells[2:]  # row 1 is the |---| separator
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        p = table.rows[0].cells[i].paragraphs[0]
        run = p.add_run(strip_inline(h))
        run.bold = True
    for row in body:
        tr = table.add_row().cells
        for i, val in enumerate(row):
            if i < len(tr):
                tr[i].paragraphs[0].add_run(strip_inline(val))


doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)

with open(SRC, encoding="utf-8") as fh:
    lines = fh.read().split("\n")

i = 0
while i < len(lines):
    line = lines[i].rstrip()

    # table block
    if line.lstrip().startswith("|"):
        block = []
        while i < len(lines) and lines[i].lstrip().startswith("|"):
            block.append(lines[i])
            i += 1
        if len(block) >= 2:
            parse_table(block)
        continue

    if not line.strip():
        i += 1
        continue

    if line.startswith("# "):
        doc.add_heading(strip_inline(line[2:]), level=0)
    elif line.startswith("## "):
        doc.add_heading(strip_inline(line[3:]), level=1)
    elif line.startswith("### "):
        doc.add_heading(strip_inline(line[4:]), level=2)
    elif line.strip() == "---":
        pass  # horizontal rule -> skip
    elif line.startswith(">"):
        p = doc.add_paragraph(style="Intense Quote")
        add_inline(p, line.lstrip("> ").strip())
    elif re.match(r"^\d+\.\s", line):
        p = doc.add_paragraph(style="List Number")
        add_inline(p, re.sub(r"^\d+\.\s", "", line))
    elif line.startswith("- "):
        p = doc.add_paragraph(style="List Bullet")
        add_inline(p, line[2:])
    else:
        p = doc.add_paragraph()
        add_inline(p, line)
    i += 1

doc.save(OUT)
print("wrote", OUT)
