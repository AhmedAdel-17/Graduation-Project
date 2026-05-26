"""
Surgical text editor for the graduation-project PPTX.

Walks all shapes (including grouped shapes), and for each text_frame whose
plain text matches an entry in REPLACEMENTS[slide_index], replaces the text
while preserving the formatting of the FIRST run.
"""
from pptx import Presentation
from copy import deepcopy

INPUT  = r"C:\Users\Seifelzeiny\Downloads\EGX Multi-Agent Stock Prediction System-2.pptx.pptx"
OUTPUT = r"C:\Users\Seifelzeiny\Downloads\EGX Multi-Agent Stock Prediction System-EDITED.pptx"

# ─────────────────────────────────────────────────────────────────────────
# Per-slide replacements. Keys are 1-indexed slide numbers.
# Each tuple is (old_text_stripped, new_text).
# Matching is done on text_frame.text.strip().
# ─────────────────────────────────────────────────────────────────────────

REPLACEMENTS: dict[int, list[tuple[str, str]]] = {
    # ──────── Slide 6 — News Analyst (relabel as 5-layer pipeline) ────────
    6: [
        ("News Analyst — Two-Layer Architecture",
         "News Analyst — 5-Layer Sentiment Pipeline"),

        # Column headers
        ("Layer 1: LLM Interpretation",
         "Layer 3: LLM Interpretation"),
        ("Layer 2: Transformer Scoring",
         "Layer 4: Transformer Scoring"),
        ("Score Blending Formula",
         "Layer 5: Blend + Coverage Penalties"),

        # Column 1 (LLM layer) — refine bullets to match code
        ("Qualitative Analysis",
         "Qualitative Reasoning"),
        ("Interprets context, sarcasm, regulatory implications",
         "Reads Arabic + English headlines; classifies sentiment with strength + confidence"),

        ("Tool-Calling Agent",
         "Skip-Tool When Prefetched"),
        ("Autonomously calls news tools when prefetched data unavailable",
         "Uses pre-fetched aggregated news (Layer 2) when available; falls back to tool calls otherwise"),

        ("Bilingual",
         "Bilingual (AR + EN)"),
        ("Interprets Arabic and English headlines natively",
         "Native Arabic (MSA + Egyptian dialect) and English interpretation"),

        ("Structured Output",
         "Structured JSON Output"),
        ("JSON with sentiment, strength, confidence, headlines",
         "sentiment, sentiment_strength, confidence_score, key_headlines[], catalysts_from_news, risks_from_news"),

        # Penalty footer note
        ("Penalty: If no transformer: combined_score = llm_score × 0.70 (30% penalty)",
         "Fallback: If transformer unavailable, combined_score = llm_score × 0.70 (30% penalty)"),

        # Column 2 (Transformer) bullets — keep mostly, tighten
        ("Specialized financial sentiment model",
         "3-class financial sentiment trained on news + tweets + Reddit"),
        ("Egyptian dialect Arabic model",
         "Egyptian-dialect Arabic; routes when language='ar'"),
        ("Multilingual fallback",
         "Cross-lingual fallback when language is mixed or unclear"),
        ("Deterministic",
         "Deterministic & Reproducible"),
        ("Reproducible scores, no prompt sensitivity",
         "No prompt sensitivity; same headlines → same scores every run"),

        # Bottom row — coverage penalties
        ("Confidence Penalties",
         "Coverage Penalties (Layer 5)"),
    ],

    # ──────── Slide 7 — News Analyst anti-feedback-loop (keep, tiny tweak) ────────
    7: [
        ("News Analyst — Anti-Feedback-Loop Design",
         "News Analyst — Layer 4 Anti-Feedback-Loop Guard"),
    ],

    # ──────── Slide 14 — Macro (refresh values + clarify) ────────
    14: [
        ("Macro Context (NOT a Standalone Agent)",
         "Macro Context — Deterministic Snapshot (NOT an Agent)"),

        # Value refreshes — match current backend / config
        # CBE Rate 27.5% kept
        ("~49.5",
         "52.85"),   # USD/EGP

        ("~$80/bbl",
         "~$109.3/bbl"),  # Brent

        ("bullish/bearish/neutral",
         "bullish/bearish/unknown"),  # EGX30 trend (unknown when yfinance fails)

        # Critical clarification text
        ("CRITICAL CLARIFICATION: Macro is NOT a Macro Analyst agent. It is a deterministic data utility fetched by DataPrefetcher.",
         "CRITICAL CLARIFICATION: Macro is NOT a Macro Analyst agent. It is a deterministic data utility called by the Data Prefetcher during Phase 0 (pre-graph). No LLM calls, no graph node, no autonomous reasoning."),

        ("Why This Matters",
         "Why This Matters for EGX"),

        ("EGX stocks are heavily influenced by macroeconomic factors: the CBE rate (27.5%) competes with equity returns, EGP depreciation affects import-heavy companies, and the IMF programme signals policy direction. Without macro context, agents would reason in a vacuum.",
         "EGX is a high-rate, currency-volatile, IMF-managed market. CBE policy rate (27.5%) competes directly with equity earnings yields; EGP at 52.85/USD reshapes every dollar-linked balance sheet; the IMF programme dictates fiscal direction. Without macro context, every agent reasons in a vacuum."),

        # Downstream consumers note
        ("Downstream Consumers: Research Manager, Trader, Risk Manager, Fundamentals Analyst (all read-only)",
         "Downstream Consumers: Research Manager, Trader, Risk Manager, Fundamentals Analyst, Bull & Bear Researchers (all read-only)"),

        ("Graceful Degradation: If yfinance fails, hard-coded defaults ensure pipeline continues (minor PIT issue in backtests)",
         "Graceful Degradation: If yfinance fails for USD/EGP or Brent, hard-coded config defaults take over so the pipeline never blocks (with a minor point-in-time impurity in backtests)."),
    ],

    # ──────── Slide 15 — Bull Researcher (tighten output schema) ────────
    15: [
        ("Bullish Thesis",
         "Strongest Long Thesis"),

        ("What It Does",
         "What It Does"),
        ("Constructs the strongest possible bullish thesis by synthesizing all analyst reports and past similar situations from memory (RAG).",
         "Synthesises all four analyst reports (Market, Fundamentals, News, Social) plus macro context, then retrieves top-3 similar long calls from bull_memory via cosine-similarity RAG to construct the strongest defensible bull thesis."),

        # Output schema fields (more accurate to actual code)
        ("STRONG_BUY / BUY / SPECULATIVE_BUY",
         "STRONG_BUY / BUY / SPECULATIVE_BUY"),
        ("HIGH / MEDIUM / LOW",
         "high / moderate / low"),
        ("target_price_range",
         "upside_scenario"),
        ("min, max, timeframe",
         "base/bull case upside %, downside risk %, basis"),
        ("SHORT_TERM / MEDIUM / LONG",
         "short_term / medium_term / long_term"),
        ("Array of catalyst objects",
         "Array of strings; explicit invalidation_conditions too"),

        # Hallucination mitigation
        ("Moderate — LLM may invent catalysts or overstate analyst signals.",
         "Moderate — LLM may invent catalysts or overstate analyst signals."),
        ("Grounded reasoning prompt, structured output, RAG with actual past cases",
         "Grounded prompt forbids inventing numbers; structured JSON enforces schema; RAG retrieves real past cases (cosine threshold = 0.30)"),

        ("Prevents speculating on excluded social data",
         "If Social Layer C emitted NO_SIGNAL, bull cannot cite social as confirmation"),

        ("Prevents speculation on excluded data",
         "Same guard — cannot fabricate social signals"),
    ],

    # ──────── Slide 16 — Bear Researcher ────────
    16: [
        ("Bearish Thesis",
         "Strongest Risk-Off Thesis"),

        ("Structurally mirrors Bull but argues the opposite side. Constructs the strongest possible bearish thesis.",
         "Structurally mirrors Bull but argues the opposite side. Pulls past failed trades and drawdowns from bear_memory to strengthen the bearish case."),

        ("Same inputs as Bull, opposite interpretation",
         "Same four analyst reports + macro; reframed to surface downside"),

        ("Past failed trades, market downturns from bear_memory",
         "Top-3 similar bearish cases retrieved via cosine-similarity RAG"),

        ("Same guard as Bull — prevents speculation",
         "Same NO_SIGNAL guard; severity tags required on every key_risk"),

        ("AVOID",
         "AVOID"),
        ("REDUCE",
         "REDUCE"),
        ("UNDERWEIGHT",
         "UNDERWEIGHT / WAIT"),

        ("Don't enter position",
         "Don't open a new position"),
        ("Trim existing position",
         "Trim an existing long position"),
        ("Hold less than benchmark",
         "Underweight vs benchmark; WAIT if data is insufficient"),

        ("Practical advice for current shareholders",
         "Mandatory field — actionable advice for current shareholders"),
        ("Array of risk objects with severity",
         "Each risk tagged severity = high/medium/low"),
        ("What would prove thesis wrong",
         "Concrete conditions that would falsify the bear case"),

        ("Regulatory Note: No short-selling recommendations — respects EGX constraints (FRA regulations)",
         "Regulatory Note: No short-selling — respects EGX/FRA constraints. SELL is exit-only; only AVOID / REDUCE / UNDERWEIGHT / WAIT are valid."),
    ],

    # ──────── Slide 17 — Research Manager (add risk_appetite) ────────
    17: [
        ("Research Manager — Judge Pattern",
         "Research Manager — Judge with Configurable Risk Appetite"),

        ("Acts as investment committee judge. Evaluates Bull vs Bear structured theses alongside macro context and portfolio position.",
         "Acts as Chief Investment Officer. Evaluates the structured Bull and Bear theses alongside macro context and current portfolio position, then issues the final BUY / SELL / HOLD decision. Now configurable per-run via a risk_appetite parameter (Conservative / Balanced / Aggressive)."),

        ("\"HOLD is a COST\"",
         "\"HOLD is a COST\""),
        ("Only valid if data insufficient OR bull/bear exactly balanced with unfavorable risk/reward. This prevents defaulting to neutral — forces a real decision.",
         "Threshold is now risk-appetite-dependent. Conservative: HOLD allowed on material deterioration even with cheap valuation. Aggressive: HOLD is the LAST RESORT — BUY when bull has any meaningful edge."),

        # 4-section output details (match actual code)
        ("1. Summary",
         "1. Strongest Bull Case"),
        ("Brief overview of bull/bear cases",
         "1-2 sentences identifying the most compelling bullish argument"),

        ("2. Analysis",
         "2. Strongest Bear Case"),
        ("Detailed evaluation of each side's strengths/weaknesses",
         "1-2 sentences identifying the most compelling bearish argument"),

        ("3. Decision",
         "3. Why One Side Wins"),
        ("BUY/SELL/HOLD with explicit reasoning",
         "2-3 sentences explaining which side has the decisive edge"),

        ("4. JSON Decision",
         "4. Decision & Plan"),
        ("action, confidence, rationale, key_factors",
         "BUY/SELL/HOLD + investment plan for Trader + JSON: {decision, confidence, rationale}"),

        ("Bull thesis (structured JSON)",
         "Bull thesis (structured JSON)"),
        ("Bear thesis (structured JSON)",
         "Bear thesis (structured JSON)"),
        ("Macro context (14+ indicators)",
         "Macro context (14+ indicators)"),
        ("Portfolio position (if any)",
         "Portfolio position + risk_appetite flag + invest_judge_memory (RAG)"),

        # Why this matters — emphasize the new risk_appetite story
        ("Why This Matters: Without the \"HOLD is a COST\" instruction, LLMs tend to default to neutral recommendations to avoid accountability. This forces a real decision.",
         "Why This Matters: In a 27.5% rate regime, conservative reasoning produced HOLD on every ticker. Exposing risk_appetite mirrors institutional reality — the same data is weighted differently by a low-vol pension fund vs. a long-biased hedge fund. Conservative reproduces prior behavior bit-for-bit; aggressive lets the bull case win on alignment alone."),
    ],
}


# ─────────────────────────────────────────────────────────────────────────
# Editor
# ─────────────────────────────────────────────────────────────────────────

def replace_in_text_frame(tf, new_text: str) -> bool:
    """
    Replace the text of `tf` with `new_text`, preserving the formatting of
    the first run (font, color, size, bold, etc.).

    Returns True on success.
    """
    if not tf.paragraphs:
        return False

    # Find the first non-empty run across all paragraphs to use as the template
    template_run = None
    for p in tf.paragraphs:
        for r in p.runs:
            if r.text:
                template_run = r
                break
        if template_run:
            break

    if template_run is None:
        # Empty frame — just set text plainly
        tf.text = new_text
        return True

    # Clone the run's properties XML
    rPr_xml = None
    if template_run._r.find('{http://schemas.openxmlformats.org/drawingml/2006/main}rPr') is not None:
        rPr_xml = deepcopy(
            template_run._r.find('{http://schemas.openxmlformats.org/drawingml/2006/main}rPr')
        )

    # First paragraph keeps formatting; subsequent paragraphs are removed
    first_para = tf.paragraphs[0]
    # Remove all runs in first paragraph except by clearing then adding new one
    # The simplest reliable approach: clear all paragraphs, add fresh
    # Approach below: preserve paragraph 0's <a:pPr>, clear runs, add new run.
    from lxml import etree
    a_ns = 'http://schemas.openxmlformats.org/drawingml/2006/main'

    # Remove all <a:r> and <a:br> children from first paragraph
    for child in list(first_para._p):
        tag = etree.QName(child).localname
        if tag in ('r', 'br', 'fld'):
            first_para._p.remove(child)

    # Remove all paragraphs after the first
    for p in list(tf.paragraphs[1:]):
        p._p.getparent().remove(p._p)

    # Add a single new run with the new text, copying rPr from template
    new_r = etree.SubElement(first_para._p, f'{{{a_ns}}}r')
    if rPr_xml is not None:
        new_r.append(rPr_xml)
    new_t = etree.SubElement(new_r, f'{{{a_ns}}}t')
    new_t.text = new_text

    return True


def walk_and_replace(shapes, mapping: dict, stats: dict):
    for sh in shapes:
        if sh.has_text_frame:
            key = sh.text_frame.text.strip()
            if key in mapping:
                new_text = mapping[key]
                if replace_in_text_frame(sh.text_frame, new_text):
                    stats["replaced"] += 1
                    stats["matched_keys"].add(key)
        # Recurse into group shapes
        if hasattr(sh, 'shape_type') and sh.shape_type == 6:
            walk_and_replace(sh.shapes, mapping, stats)


def main():
    p = Presentation(INPUT)
    total_stats = {"replaced": 0, "missed": []}

    for slide_idx, repls in REPLACEMENTS.items():
        slide = p.slides[slide_idx - 1]
        mapping = {old: new for old, new in repls}
        stats = {"replaced": 0, "matched_keys": set()}
        walk_and_replace(slide.shapes, mapping, stats)
        missed = [k for k in mapping if k not in stats["matched_keys"]]
        total_stats["replaced"] += stats["replaced"]
        if missed:
            print(f"Slide {slide_idx}: replaced {stats['replaced']} / {len(repls)}  | MISSED: {missed}")
        else:
            print(f"Slide {slide_idx}: replaced {stats['replaced']} / {len(repls)}  ✓")

    p.save(OUTPUT)
    print(f"\nSaved to: {OUTPUT}")
    print(f"Total replacements: {total_stats['replaced']}")


if __name__ == "__main__":
    main()
